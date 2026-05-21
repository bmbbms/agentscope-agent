# -*- coding: utf-8 -*-
"""
数据处理模块
负责读取Excel数据、计算指标、生成看板数据
"""

import json
from typing import Dict, List, Any, Optional
from openpyxl import load_workbook
from datetime import datetime


class DataProcessor:
    """数据处理类"""

    def __init__(self, data_file: str):
        """
        初始化数据处理器

        Args:
            data_file: 数据文件路径
        """
        self.data_file = data_file
        self.wb = load_workbook(data_file, data_only=True)

    def get_sheets_data(self) -> Dict[str, List[Dict]]:
        """
        获取所有Sheet的数据

        Returns:
            按Sheet名称组织的数据字典
        """
        result = {}

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            headers = [cell.value for cell in ws[1] if cell.value]

            data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = {}
                for i, header in enumerate(headers):
                    if i < len(row):
                        value = row[i]
                        # 处理None值
                        if value is None:
                            value = 0 if any(kw in str(header) for kw in ['金额', '笔数', '商户', '数', '万', '元', '台']) else ''
                        row_dict[header] = value
                if any(v for v in row_dict.values() if v):  # 过滤空行
                    data.append(row_dict)

            result[sheet_name] = data

        return result

    def calculate_small_pos_dashboard(self) -> Dict[str, Any]:
        """
        计算立刷产品看板数据（小POS，product_type=5）

        Returns:
            看板数据
        """
        data = self.get_sheets_data()

        # 1. 月总体数据
        monthly_data = data.get('立刷月总体', [])
        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        # 计算环比（原看板只有4个核心指标，不含交易笔数）
        mom_change = {}
        if latest_monthly and previous_monthly:
            for key in ['交易总金额（万元）', '活跃商户数', '当月笔均金额（元）', '新增商户数']:
                if key in latest_monthly and key in previous_monthly:
                    current = float(latest_monthly.get(key, 0) or 0)
                    previous = float(previous_monthly.get(key, 0) or 0)
                    if previous > 0:
                        mom_change[key] = round((current - previous) / previous * 100, 2)

        # 2. 产品新增金额趋势
        product_amount = data.get('立刷产品新增金额', [])

        # 3. 产品新增终端趋势
        product_terminal = data.get('立刷产品新增终端', [])

        # 4. 服务商排名（下降TOP10）
        agent_decline_trade = data.get('立刷服务商新增交易下降TOP10', [])
        agent_decline_merchant = data.get('立刷服务商新增商户下降排名', [])

        return {
            'title': '立刷产品看板（小POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': latest_monthly.get('月份', ''),
            # 核心指标（4个，与原看板一致）
            'summary': {
                '交易总金额': {
                    'value': float(latest_monthly.get('交易总金额（万元）', 0) or 0),
                    'unit': '亿元',
                    'display_value': round(float(latest_monthly.get('交易总金额（万元）', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('交易总金额（万元）', 0)
                },
                '活跃商户数': {
                    'value': float(latest_monthly.get('活跃商户数', 0) or 0),
                    'unit': '万户',
                    'display_value': round(float(latest_monthly.get('活跃商户数', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('活跃商户数', 0)
                },
                '笔均金额': {
                    'value': float(latest_monthly.get('当月笔均金额（元）', 0) or 0),
                    'unit': '元',
                    'display_value': int(float(latest_monthly.get('当月笔均金额（元）', 0) or 0)),
                    'mom': mom_change.get('当月笔均金额（元）', 0)
                },
                '新增商户数': {
                    'value': float(latest_monthly.get('新增商户数', 0) or 0),
                    'unit': '万户',
                    'display_value': round(float(latest_monthly.get('新增商户数', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('新增商户数', 0)
                }
            },
            'monthly_trend': monthly_data[-12:] if len(monthly_data) >= 12 else monthly_data,  # 最近12个月
            'product_amount_trend': product_amount,
            'product_terminal_trend': product_terminal,
            'agent_decline_trade': agent_decline_trade[:10],
            'agent_decline_merchant': agent_decline_merchant[:10]
        }

    def calculate_large_pos_dashboard(self) -> Dict[str, Any]:
        """
        计算商户收款看板数据（大POS，product_type=2）

        Returns:
            看板数据
        """
        data = self.get_sheets_data()

        # 1. 月总体数据
        monthly_data = data.get('大POS月总体', [])
        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        # 计算环比（原看板只有4个核心指标，不含交易笔数）
        mom_change = {}
        if latest_monthly and previous_monthly:
            for key in ['交易总金额（万元）', '活跃用户数', '当月笔均金额（元）', '新增商户数']:
                if key in latest_monthly and key in previous_monthly:
                    current = float(latest_monthly.get(key, 0) or 0)
                    previous = float(previous_monthly.get(key, 0) or 0)
                    if previous > 0:
                        mom_change[key] = round((current - previous) / previous * 100, 2)

        # 2. 近两年运营数据
        two_year_data = data.get('大POS近两年运营数据', [])

        # 3. 分公司排名（下降TOP10）
        branch_decline_trade = data.get('大POS同期新增交易下降TOP10', [])
        branch_decline_merchant = data.get('大POS同期新增商户下降TOP10', [])

        # 4. 商户排名
        merchant_decline = data.get('大POS商户新增下降TOP10', [])
        merchant_increase = data.get('大POS商户新增上涨TOP10', [])

        return {
            'title': '商户收款看板（大POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': latest_monthly.get('月份', ''),
            # 核心指标（4个，与原看板一致）
            'summary': {
                '交易总金额': {
                    'value': float(latest_monthly.get('交易总金额（万元）', 0) or 0),
                    'unit': '亿元',
                    'display_value': round(float(latest_monthly.get('交易总金额（万元）', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('交易总金额（万元）', 0)
                },
                '活跃用户数': {
                    'value': float(latest_monthly.get('活跃用户数', 0) or 0),
                    'unit': '万户',
                    'display_value': round(float(latest_monthly.get('活跃用户数', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('活跃用户数', 0)
                },
                '笔均金额': {
                    'value': float(latest_monthly.get('当月笔均金额（元）', 0) or 0),
                    'unit': '元',
                    'display_value': int(float(latest_monthly.get('当月笔均金额（元）', 0) or 0)),
                    'mom': mom_change.get('当月笔均金额（元）', 0)
                },
                '新增商户数': {
                    'value': float(latest_monthly.get('新增商户数', 0) or 0),
                    'unit': '万户',
                    'display_value': round(float(latest_monthly.get('新增商户数', 0) or 0) / 10000, 2),
                    'mom': mom_change.get('新增商户数', 0)
                }
            },
            'monthly_trend': monthly_data,
            'two_year_trend': two_year_data[-24:] if len(two_year_data) >= 24 else two_year_data,
            'branch_decline_trade': branch_decline_trade[:10],
            'branch_decline_merchant': branch_decline_merchant[:10],
            'merchant_decline': merchant_decline[:10],
            'merchant_increase': merchant_increase[:10],
            # 细分业务
            'business_breakdown': {
                '外接码付': {
                    '金额': float(latest_monthly.get('外接码付(咕朵云)交易金额/万元', 0) or 0),
                    '笔数': float(latest_monthly.get('外接码付(咕朵云)交易笔数/万笔', 0) or 0)
                },
                '理财POS': {
                    '金额': float(latest_monthly.get('理财POS交易金额/万元', 0) or 0),
                    '笔数': float(latest_monthly.get('理财POS交易笔数/万笔', 0) or 0)
                },
                '传统POS': {
                    '金额': float(latest_monthly.get('传统POS交易金额/万元', 0) or 0),
                    '笔数': float(latest_monthly.get('传统POS交易笔数/万笔', 0) or 0)
                }
            }
        }

    def close(self):
        """关闭工作簿"""
        self.wb.close()


def format_number(value: Any, decimal: int = 2) -> str:
    """
    格式化数字显示

    Args:
        value: 数值
        decimal: 小数位数

    Returns:
        格式化后的字符串
    """
    try:
        num = float(value)
        if num == 0:
            return '-'
        return f'{num:,.{decimal}f}'
    except:
        return str(value)


def format_percent(value: Any) -> str:
    """
    格式化百分比显示

    Args:
        value: 数值

    Returns:
        格式化后的字符串
    """
    try:
        num = float(value)
        if num == 0:
            return '-'
        sign = '+' if num > 0 else ''
        return f'{sign}{num:.2f}%'
    except:
        return str(value)


def get_trend_class(value: Any) -> str:
    """
    获取趋势样式类

    Args:
        value: 数值

    Returns:
        CSS类名
    """
    try:
        num = float(value)
        if num > 0:
            return 'trend-up'
        elif num < 0:
            return 'trend-down'
        return ''
    except:
        return ''


if __name__ == '__main__':
    # 测试数据处理
    processor = DataProcessor('../收单运营月报看板导出数据.xlsx')

    print('=== 立刷产品看板数据 ===')
    small_pos = processor.calculate_small_pos_dashboard()
    print(f'最新月份: {small_pos["latest_month"]}')
    print(f'汇总数据: {small_pos["summary"]}')

    print()
    print('=== 商户收款看板数据 ===')
    large_pos = processor.calculate_large_pos_dashboard()
    print(f'最新月份: {large_pos["latest_month"]}')
    print(f'汇总数据: {large_pos["summary"]}')

    processor.close()