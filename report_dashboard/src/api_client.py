# -*- coding: utf-8 -*-
"""
数据平台API接口对接模块
支持从Superset BI获取真实数据
"""

import os
import sys
from typing import Dict, List, Optional, Any
from datetime import datetime

# superset_client路径（仅在api模式时导入）
_superset_path = os.path.join(os.path.dirname(__file__), 'superset_client')

from dashboard_queries import (
    get_date_params,
    translate_fields,
    format_number,
    SQL_LARGE_POS_MONTHLY,
    SQL_LARGE_POS_TWO_YEAR,
    SQL_LARGE_POS_BRANCH_DECLINE,
    SQL_LARGE_POS_BRANCH_MERCHANT_DECLINE,
    SQL_LARGE_POS_MERCHANT_DECLINE,
    SQL_LARGE_POS_MERCHANT_INCREASE,
    SQL_LARGE_POS_BRANCH_INCREASE,
    SQL_LARGE_POS_BRANCH_MERCHANT_INCREASE,
    SQL_SMALL_POS_MONTHLY,
    SQL_SMALL_POS_TWO_YEAR,
    SQL_SMALL_POS_AGENT_DECLINE,
    SQL_SMALL_POS_AGENT_MERCHANT_DECLINE,
    SQL_SMALL_POS_AGENT_INCREASE,
    SQL_SMALL_POS_AGENT_MERCHANT_INCREASE,
    SQL_BUSINESS_TYPE_AMOUNT,
    SQL_BUSINESS_TYPE_COUNT,
    SQL_SMALL_POS_PRODUCT_AMOUNT,
    SQL_SMALL_POS_PRODUCT_TERMINAL,
    SQL_SMALL_POS_TERM_TREND,
)


class SupersetDataFetcher:
    """Superset数据获取器"""

    def __init__(self, base_url: str = None, username: str = None, password: str = None):
        """
        初始化数据获取器

        Args:
            base_url: Superset服务地址
            username: 用户名
            password: 密码
        """
        # 仅在api模式下导入superset_client（避免pandas/numpy兼容性问题）
        if _superset_path not in sys.path:
            sys.path.insert(0, _superset_path)
        from superset_client import SupersetClient, get_config

        self._SupersetClient = SupersetClient
        config = get_config()
        self.base_url = base_url or config.get('url', '') or os.environ.get('SUPERSET_URL', '')
        self.username = username or config.get('username', '') or os.environ.get('SUPERSET_USERNAME', '')
        self.password = password or config.get('password', '') or os.environ.get('SUPERSET_PASSWORD', '')
        self.client = None
        self._connected = False

    def connect(self) -> bool:
        """连接Superset"""
        if not self.base_url or not self.username or not self.password:
            print("错误：缺少Superset连接配置")
            print("请设置环境变量：")
            print("  SUPERSET_URL - Superset服务地址")
            print("  SUPERSET_USERNAME - 用户名")
            print("  SUPERSET_PASSWORD - 密码")
            return False

        self.client = self._SupersetClient(self.base_url, self.username, self.password)
        self._connected = self.client.login()
        return self._connected

    def execute_query(self, sql: str, database_name: str = None) -> List[Dict]:
        """
        执行SQL查询

        Args:
            sql: SQL语句
            database_name: 数据库名称

        Returns:
            查询结果列表（字段名已转换为中文）
        """
        if not self._connected:
            if not self.connect():
                return []

        result = self.client.execute_sql(sql, database_name=database_name)
        if result is not None and len(result) > 0:
            # 将英文字段名转换为中文
            return translate_fields(result)
        return []

    def _get_latest_data_month(self) -> str:
        """
        获取数据可用的最新月份

        Returns:
            最新月份字符串，格式YYYYMM
        """
        # 确保已连接
        if not self._connected:
            if not self.connect():
                return None

        sql = "SELECT MAX(dt) as max_dt FROM edw.agt_trd_stat_1m"
        result = self.client.execute_sql(sql)
        if result and len(result) > 0:
            max_dt = result[0].get('max_dt')
            if max_dt:
                return str(max_dt)
        return None

    def fetch_large_pos_dashboard(self, stat_month: str = None) -> Dict[str, Any]:
        """
        获取商户收款看板数据（大POS）

        Args:
            stat_month: 统计月份，格式YYYYMM

        Returns:
            看板数据字典
        """
        # 如果未指定统计月份，自动检测最新数据月份
        if stat_month is None:
            stat_month = self._get_latest_data_month()
            if stat_month:
                print(f"  - 自动检测到最新数据月份: {stat_month}")

        params = get_date_params(stat_month)

        # 1. 月度总体数据
        monthly_data = self.execute_query(
            SQL_LARGE_POS_MONTHLY.format(start_dt=params['start_dt'], end_dt=params['end_dt'])
        )

        # 2. 近两年运营数据
        two_year_data = self.execute_query(
            SQL_LARGE_POS_TWO_YEAR.format(start_dt=params['start_dt'], end_dt=params['end_dt'])
        )

        # 3. 分公司交易下降TOP10
        branch_decline = self.execute_query(
            SQL_LARGE_POS_BRANCH_DECLINE.format(
                current_dt=params['current_dt'],
                previous_dt=params['previous_dt']
            )
        )

        # 4. 分公司商户下降TOP10
        branch_merchant_decline = self.execute_query(
            SQL_LARGE_POS_BRANCH_MERCHANT_DECLINE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 5. 商户交易下降TOP10
        merchant_decline = self.execute_query(
            SQL_LARGE_POS_MERCHANT_DECLINE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 6. 商户交易上涨TOP10
        merchant_increase = self.execute_query(
            SQL_LARGE_POS_MERCHANT_INCREASE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 7. 分公司交易上涨TOP10
        branch_increase = self.execute_query(
            SQL_LARGE_POS_BRANCH_INCREASE.format(
                current_dt=params['current_dt'],
                previous_dt=params['previous_dt']
            )
        )

        # 8. 分公司商户上涨TOP10
        branch_merchant_increase = self.execute_query(
            SQL_LARGE_POS_BRANCH_MERCHANT_INCREASE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 9. 业务类型趋势（近13个月）
        business_amount = self.execute_query(
            SQL_BUSINESS_TYPE_AMOUNT.format(start_dt=params['start_dt_1year'], end_dt=params['end_dt'])
        )
        business_count = self.execute_query(
            SQL_BUSINESS_TYPE_COUNT.format(start_dt=params['start_dt_1year'], end_dt=params['end_dt'])
        )

        # 计算核心指标
        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        summary = self._calculate_summary(latest_monthly, previous_monthly, '活跃商户数')

        return {
            'title': '商户收款看板（大POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': params['stat_month'],
            'summary': summary,
            'monthly_trend': monthly_data[-12:] if monthly_data else [],
            'two_year_trend': two_year_data,
            'branch_decline_trade': branch_decline,
            'branch_decline_merchant': branch_merchant_decline,
            'branch_increase_trade': branch_increase,
            'branch_increase_merchant': branch_merchant_increase,
            'merchant_decline': merchant_decline,
            'merchant_increase': merchant_increase,
            'business_amount_trend': business_amount[-13:] if business_amount else [],
            'business_count_trend': business_count[-13:] if business_count else [],
        }

    def fetch_small_pos_dashboard(self, stat_month: str = None) -> Dict[str, Any]:
        """
        获取立刷产品看板数据（小POS）

        Args:
            stat_month: 统计月份，格式YYYYMM

        Returns:
            看板数据字典
        """
        # 如果未指定统计月份，自动检测最新数据月份
        if stat_month is None:
            stat_month = self._get_latest_data_month()

        params = get_date_params(stat_month)

        # 1. 月度总体数据
        monthly_data = self.execute_query(
            SQL_SMALL_POS_MONTHLY.format(start_dt=params['start_dt'], end_dt=params['end_dt'])
        )

        # 2. 近两年运营数据
        two_year_data = self.execute_query(
            SQL_SMALL_POS_TWO_YEAR.format(start_dt=params['start_dt'], end_dt=params['end_dt'])
        )

        # 3. 服务商交易下降TOP10
        agent_decline = self.execute_query(
            SQL_SMALL_POS_AGENT_DECLINE.format(
                current_dt=params['current_dt'],
                previous_dt=params['previous_dt']
            )
        )

        # 4. 服务商商户下降TOP10
        agent_merchant_decline = self.execute_query(
            SQL_SMALL_POS_AGENT_MERCHANT_DECLINE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 5. 服务商交易上涨TOP10
        agent_increase = self.execute_query(
            SQL_SMALL_POS_AGENT_INCREASE.format(
                current_dt=params['current_dt'],
                previous_dt=params['previous_dt']
            )
        )

        # 6. 服务商商户上涨TOP10
        agent_merchant_increase = self.execute_query(
            SQL_SMALL_POS_AGENT_MERCHANT_INCREASE.format(
                current_month=params['current_month'],
                previous_month=params['previous_month']
            )
        )

        # 7. 各产品交易金额趋势（近13个月）
        product_amount = self.execute_query(
            SQL_SMALL_POS_PRODUCT_AMOUNT.format(start_dt=params['start_dt_1year'], end_dt=params['end_dt'])
        )

        # 8. 各产品新增终端趋势（近13个月）
        product_terminal = self.execute_query(
            SQL_SMALL_POS_PRODUCT_TERMINAL.format(start_dt=params['start_dt_1year'], end_dt=params['end_dt'])
        )

        # 9. 终端绑定、激活、活跃趋势（近25个月）
        term_trend = self.execute_query(
            SQL_SMALL_POS_TERM_TREND.format(start_dt=params['start_dt'], end_dt=params['end_dt'])
        )

        # 计算核心指标
        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        summary = self._calculate_summary(latest_monthly, previous_monthly, '活跃商户数')

        return {
            'title': '立刷产品看板（小POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': params['stat_month'],
            'summary': summary,
            'monthly_trend': monthly_data,  # 近25个月数据（SQL已返回正确范围）
            'two_year_trend': two_year_data,
            'agent_decline_trade': agent_decline,
            'agent_decline_merchant': agent_merchant_decline,
            'agent_increase_trade': agent_increase,
            'agent_increase_merchant': agent_merchant_increase,
            'product_amount_trend': product_amount[-13:] if product_amount else [],
            'product_terminal_trend': product_terminal[-13:] if product_terminal else [],
            'term_trend': term_trend,  # 终端绑定、激活、活跃趋势
        }

    def _calculate_summary(self, latest: Dict, previous: Dict, active_field: str) -> Dict:
        """计算核心指标汇总（计算所有可用指标，由模板控制显示）"""
        summary = {}

        # 交易总金额（单位：万元）
        current_amount = self._safe_float(latest.get('交易总金额（万元）', 0))
        previous_amount = self._safe_float(previous.get('交易总金额（万元）', 0))
        summary['交易总金额'] = {
            'value': current_amount,
            'unit': '万元',
            'display_value': round(current_amount) if current_amount else 0,
            'mom': self._calc_mom(current_amount, previous_amount)
        }

        # 新增商户数（单位：户）
        current_mer = self._safe_float(latest.get('新增商户数', 0))
        previous_mer = self._safe_float(previous.get('新增商户数', 0))
        summary['新增商户数'] = {
            'value': current_mer,
            'unit': '户',
            'display_value': round(current_mer) if current_mer else 0,
            'mom': self._calc_mom(current_mer, previous_mer)
        }

        # 活跃商户数（单位：户）
        current_active = self._safe_float(latest.get('活跃商户数', 0))
        previous_active = self._safe_float(previous.get('活跃商户数', 0))
        summary['活跃商户数'] = {
            'value': current_active,
            'unit': '户',
            'display_value': round(current_active) if current_active else 0,
            'mom': self._calc_mom(current_active, previous_active)
        }

        # 交易总笔数（单位：万笔）- 可选指标
        current_count = self._safe_float(latest.get('交易总笔数（万笔）', 0))
        previous_count = self._safe_float(previous.get('交易总笔数（万笔）', 0))
        summary['交易总笔数'] = {
            'value': current_count,
            'unit': '万笔',
            'display_value': round(current_count) if current_count else 0,
            'mom': self._calc_mom(current_count, previous_count)
        }

        return summary

    def _safe_float(self, value) -> float:
        """安全转换为浮点数"""
        try:
            return float(value) if value else 0
        except:
            return 0

    def _calc_mom(self, current: float, previous: float) -> float:
        """计算环比"""
        if previous > 0:
            return round((current - previous) / previous * 100, 2)
        return 0


class MockDataFetcher:
    """模拟数据获取器（从本地Excel读取）"""

    def __init__(self, data_file: str):
        self.data_file = data_file
        self._load_data()

    def _load_data(self):
        """加载本地数据"""
        from openpyxl import load_workbook
        self.wb = load_workbook(self.data_file, data_only=True)
        self.sheets_data = {}
        self.sheets_raw = {}  # 存储原始数据（含重复列名）

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            headers = [cell.value for cell in ws[1] if cell.value]
            data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = {}
                for i, header in enumerate(headers):
                    if i < len(row):
                        value = row[i]
                        if value is None:
                            value = 0 if any(kw in str(header) for kw in ['金额', '笔数', '商户', '数', '万', '元', '台']) else ''
                        row_dict[header] = value
                if any(v for v in row_dict.values() if v):
                    data.append(row_dict)
            self.sheets_data[sheet_name] = data
            # 同时存储原始行数据（按索引访问）
            self.sheets_raw[sheet_name] = {
                'headers': headers,
                'rows': list(ws.iter_rows(min_row=2, values_only=True))
            }

    def fetch_large_pos_dashboard(self, stat_month: str = None) -> Dict[str, Any]:
        """获取大POS看板数据"""
        monthly_data = self.sheets_data.get('大POS月总体', [])
        two_year_data = self.sheets_data.get('大POS近两年运营数据', [])

        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        # 计算环比
        mom_change = {}
        if latest_monthly and previous_monthly:
            for key in ['交易总金额（万元）', '活跃用户数', '当月笔均金额（元）', '新增商户数']:
                if key in latest_monthly and key in previous_monthly:
                    current = float(latest_monthly.get(key, 0) or 0)
                    previous = float(previous_monthly.get(key, 0) or 0)
                    if previous > 0:
                        mom_change[key] = round((current - previous) / previous * 100, 2)

        return {
            'title': '商户收款看板（大POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': latest_monthly.get('月份', ''),
            'summary': {
                '交易总金额': {
                    'value': float(latest_monthly.get('交易总金额（万元）', 0) or 0),
                    'unit': '万元',
                    'display_value': round(float(latest_monthly.get('交易总金额（万元）', 0) or 0), 0),
                    'mom': mom_change.get('交易总金额（万元）', 0)
                },
                '新增商户数': {
                    'value': float(latest_monthly.get('新增商户数', 0) or 0),
                    'unit': '户',
                    'display_value': round(float(latest_monthly.get('新增商户数', 0) or 0), 0),
                    'mom': mom_change.get('新增商户数', 0)
                },
                '活跃商户数': {
                    'value': float(latest_monthly.get('活跃用户数', 0) or 0),
                    'unit': '户',
                    'display_value': round(float(latest_monthly.get('活跃用户数', 0) or 0), 0),
                    'mom': mom_change.get('活跃用户数', 0)
                },
                '交易总笔数': {
                    'value': float(latest_monthly.get('交易总笔数（万笔）', 0) or 0),
                    'unit': '万笔',
                    'display_value': round(float(latest_monthly.get('交易总笔数（万笔）', 0) or 0), 0),
                    'mom': mom_change.get('交易总笔数（万笔）', 0)
                }
            },
            'monthly_trend': monthly_data,
            'two_year_trend': self._normalize_two_year_data(two_year_data),
            'branch_decline_trade': self._normalize_branch_trade_data(self.sheets_data.get('大POS同期新增交易下降TOP10', [])),
            'branch_decline_merchant': self._normalize_branch_merchant_data(self.sheets_data.get('大POS同期新增商户下降TOP10', [])),
            'branch_increase_trade': [],
            'branch_increase_merchant': [],
            'merchant_decline': self._normalize_merchant_data(self.sheets_data.get('大POS商户新增下降TOP10', [])),
            'merchant_increase': self._normalize_merchant_data(self.sheets_data.get('大POS商户新增上涨TOP10', [])),
            'business_amount_trend': self._extract_business_amount_trend(monthly_data),
            'business_count_trend': self._extract_business_count_trend(monthly_data),
        }

    def fetch_small_pos_dashboard(self, stat_month: str = None) -> Dict[str, Any]:
        """获取小POS看板数据"""
        monthly_data = self.sheets_data.get('立刷月总体', [])
        product_amount = self.sheets_data.get('立刷产品新增金额', [])
        product_terminal = self.sheets_data.get('立刷产品新增终端', [])

        latest_monthly = monthly_data[-1] if monthly_data else {}
        previous_monthly = monthly_data[-2] if len(monthly_data) >= 2 else {}

        # 计算环比
        mom_change = {}
        if latest_monthly and previous_monthly:
            for key in ['交易总金额（万元）', '活跃商户数', '当月笔均金额（元）', '新增商户数']:
                if key in latest_monthly and key in previous_monthly:
                    current = float(latest_monthly.get(key, 0) or 0)
                    previous = float(previous_monthly.get(key, 0) or 0)
                    if previous > 0:
                        mom_change[key] = round((current - previous) / previous * 100, 2)

        return {
            'title': '立刷产品看板（小POS）',
            'update_time': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'latest_month': latest_monthly.get('月份', ''),
            'summary': {
                '交易总金额': {
                    'value': float(latest_monthly.get('交易总金额（万元）', 0) or 0),
                    'unit': '万元',
                    'display_value': round(float(latest_monthly.get('交易总金额（万元）', 0) or 0), 0),
                    'mom': mom_change.get('交易总金额（万元）', 0)
                },
                '新增商户数': {
                    'value': float(latest_monthly.get('新增商户数', 0) or 0),
                    'unit': '户',
                    'display_value': round(float(latest_monthly.get('新增商户数', 0) or 0), 0),
                    'mom': mom_change.get('新增商户数', 0)
                },
                '活跃商户数': {
                    'value': float(latest_monthly.get('活跃商户数', 0) or 0),
                    'unit': '户',
                    'display_value': round(float(latest_monthly.get('活跃商户数', 0) or 0), 0),
                    'mom': mom_change.get('活跃商户数', 0)
                }
            },
            'monthly_trend': monthly_data,
            'two_year_trend': self._normalize_small_pos_two_year(monthly_data),
            'term_trend': self._normalize_term_trend(monthly_data),
            'product_amount_trend': self._normalize_product_amount(product_amount),
            'product_terminal_trend': self._normalize_product_terminal(product_terminal),
            'agent_decline_trade': self._normalize_agent_trade(self.sheets_data.get('立刷服务商新增交易下降TOP10', [])),
            'agent_decline_merchant': self._normalize_agent_merchant(self.sheets_data.get('立刷服务商新增商户下降排名', [])),
        }

    def _extract_business_amount_trend(self, monthly_data: list) -> list:
        """从月度数据中提取各业务类型交易金额趋势"""
        result = []
        for row in monthly_data:
            item = {
                '月份': row.get('月份', ''),
                # 使用html_generator期望的键名
                '外接码付(咕哚云)交易金额（万元）': float(row.get('外接码付(咕朵云)交易金额/万元', 0) or 0),
                '理财POS交易金额（万元）': float(row.get('理财POS交易金额/万元', 0) or 0),
                '传统POS交易金额（万元）': float(row.get('传统POS交易金额/万元', 0) or 0),
            }
            result.append(item)
        return result[-13:] if result else []  # 返回近13个月

    def _extract_business_count_trend(self, monthly_data: list) -> list:
        """从月度数据中提取各业务类型交易笔数趋势"""
        result = []
        for row in monthly_data:
            item = {
                '月份': row.get('月份', ''),
                # 使用html_generator期望的键名
                '外接码付(咕哚云)交易笔数（万笔）': float(row.get('外接码付(咕朵云)交易笔数/万笔', 0) or 0),
                '理财POS交易笔数（万笔）': float(row.get('理财POS交易笔数/万笔', 0) or 0),
                '传统POS交易笔数（万笔）': float(row.get('传统POS交易笔数/万笔', 0) or 0),
            }
            result.append(item)
        return result[-13:] if result else []  # 返回近13个月

    def _normalize_two_year_data(self, data: list) -> list:
        """规范化大POS近两年运营数据字段名"""
        result = []
        for row in data:
            item = {
                '月份': row.get('时间', ''),
                '交易总金额（万元）': float(row.get('交易总金额（万元）', 0) or 0),
                '新增商户数': float(row.get('新增商户数', 0) or 0),
                '活跃商户数': float(row.get('活跃商户数', 0) or 0),
            }
            result.append(item)
        return result[-24:] if result else []

    def _normalize_branch_trade_data(self, data: list) -> list:
        """规范化分公司交易下降TOP10数据"""
        result = []
        for row in data[:10]:
            mom = float(row.get('环比', 0) or 0)
            item = {
                '归属分公司': row.get('归属分公司', ''),
                '上月交易额（万元）': float(row.get('上月交易额/万元', 0) or 0),
                '本月交易额（万元）': float(row.get('本月交易额/万元', 0) or 0),
                '环比': round(mom * 100, 2),  # 转换为百分比
            }
            result.append(item)
        return result

    def _normalize_branch_merchant_data(self, data: list) -> list:
        """规范化分公司商户下降TOP10数据"""
        result = []
        for row in data[:10]:
            mom = float(row.get('环比', 0) or 0)
            item = {
                '归属分公司': row.get('归属分公司', ''),
                '上月商户数': float(row.get('上月新增商户数', 0) or 0),
                '本月商户数': float(row.get('本月新增商户数', 0) or 0),
                '环比': round(mom * 100, 2),  # 转换为百分比
            }
            result.append(item)
        return result

    def _normalize_merchant_data(self, data: list) -> list:
        """规范化商户TOP10数据"""
        result = []
        for row in data[:10]:
            mom = float(row.get('环比', 0) or 0)
            item = {
                '子商户名': row.get('子商户名', ''),
                '归属分公司': row.get('归属分公司', ''),
                '上月交易额（万元）': float(row.get('上月交易额/万元', 0) or 0),
                '本月交易额（万元）': float(row.get('本月交易额/万元', 0) or 0),
                '环比': round(mom * 100, 2),  # 转换为百分比
            }
            result.append(item)
        return result

    def _normalize_small_pos_two_year(self, monthly_data: list) -> list:
        """从立刷月总体数据生成近两年趋势"""
        result = []
        for row in monthly_data[-24:]:
            item = {
                '月份': row.get('月份', ''),
                '新增绑定': float(row.get('新增绑定', 0) or 0),
                '一阶段激活数': float(row.get('一阶段激活数', 0) or 0),
                '活跃商户数': float(row.get('活跃商户数', 0) or 0),
            }
            result.append(item)
        return result

    def _normalize_term_trend(self, monthly_data: list) -> list:
        """规范化终端绑定、激活、活跃趋势数据"""
        result = []
        for row in monthly_data[-24:]:
            item = {
                '月份': row.get('月份', ''),
                '新增终端绑定（台）': float(row.get('新增绑定', 0) or 0),
                '一阶段激活数（台）': float(row.get('一阶段激活数', 0) or 0),
                '活跃商户数（台）': float(row.get('活跃商户数', 0) or 0),
            }
            result.append(item)
        return result

    def _normalize_product_amount(self, data: list) -> list:
        """规范化立刷产品金额趋势数据"""
        result = []
        for row in data:
            item = {
                '月份': row.get('日期', ''),
                '微电签交易金额（亿元）': float(row.get('微电签-金额（亿元）', 0) or 0),
                '微智能交易金额（亿元）': float(row.get('微智能-金额（亿元）', 0) or 0),
                '立刷畅享版交易金额（亿元）': float(row.get('立刷畅享版（亿元）', 0) or 0),
                '骐骥版交易金额（亿元）': float(row.get('骐骥版交易金额（亿元）', 0) or 0),
                '老电签交易金额（亿元）': float(row.get('老电签-金额（亿元）', 0) or 0),
                '上网宝交易金额（亿元）': float(row.get('上网宝-金额（亿元）', 0) or 0),
                '立刷小蓝牙交易金额（亿元）': float(row.get('立刷小蓝牙-金额（亿元）', 0) or 0),
            }
            result.append(item)
        return result[-13:] if result else []

    def _normalize_product_terminal(self, data: list) -> list:
        """规范化立刷产品终端趋势数据"""
        result = []
        for row in data:
            item = {
                '月份': row.get('日期', ''),
                '微电签新增终端（万台）': float(row.get('微电签-新增终端（万台）', 0) or 0),
                '微智能新增终端（万台）': float(row.get('微智能-新增终端（万台）', 0) or 0),
                '立刷畅享版新增终端（万台）': float(row.get('立刷畅享版-新增终端（万台）', 0) or 0),
                '骐骥版新增终端（万台）': float(row.get('骐骥版新增终端（万台）', 0) or 0),
                '老电签新增终端（万台）': float(row.get('老电签-新增终端（万台）', 0) or 0),
                '上网宝新增终端（万台）': float(row.get('上网宝-新增终端（万台）', 0) or 0),
                '立刷小蓝牙新增终端（万台）': float(row.get('立刷小蓝牙-新增终端（万台）', 0) or 0),
            }
            result.append(item)
        return result[-13:] if result else []

    def _normalize_agent_trade(self, data: list) -> list:
        """规范化服务商交易下降TOP10数据"""
        # 使用原始数据按索引获取（避免重复列名问题）
        raw = self.sheets_raw.get('立刷服务商新增交易下降TOP10', {})
        rows = raw.get('rows', [])

        result = []
        for i, row in enumerate(rows[:10]):
            if len(row) >= 5:
                mom = float(row[4] if row[4] else 0)
                item = {
                    '一级代理商名称': row[1] if len(row) > 1 else '',
                    '上月交易额（万元）': float(row[2] if row[2] else 0),
                    '本月交易额（万元）': float(row[3] if row[3] else 0),
                    '环比': round(mom * 100, 2),  # 转换为百分比
                }
                result.append(item)
        return result

    def _normalize_agent_merchant(self, data: list) -> list:
        """规范化服务商商户下降TOP10数据"""
        # 使用原始数据按索引获取
        raw = self.sheets_raw.get('立刷服务商新增商户下降排名', {})
        rows = raw.get('rows', [])

        result = []
        for row in rows[:10]:
            if len(row) >= 5:
                mom = float(row[4] if row[4] else 0)
                item = {
                    '一级代理商名称': row[1] if len(row) > 1 else '',
                    '上月商户数': float(row[2] if row[2] else 0),
                    '本月商户数': float(row[3] if row[3] else 0),
                    '环比': round(mom * 100, 2),  # 转换为百分比
                }
                result.append(item)
        return result

    def close(self):
        """关闭数据源"""
        if self.wb:
            self.wb.close()


def create_data_fetcher(config: Dict):
    """
    创建数据获取器

    Args:
        config: 配置字典，包含:
            - mode: 'api' 或 'mock'
            - data_file: 本地数据文件路径（mock模式）
            - base_url, username, password: Superset配置（api模式）

    Returns:
        数据获取器实例
    """
    mode = config.get('mode', 'mock')

    if mode == 'api':
        return SupersetDataFetcher(
            base_url=config.get('base_url'),
            username=config.get('username'),
            password=config.get('password')
        )
    else:
        return MockDataFetcher(data_file=config['data_file'])